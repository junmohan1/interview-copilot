"""文件解析：PDF / DOCX / TXT / MD / XLSX"""
from pathlib import Path


def parse_file(path: str) -> str:
    p = Path(path)
    suffix = p.suffix.lower()

    if suffix == ".pdf":
        return _parse_pdf(p)
    elif suffix == ".docx":
        return _parse_docx(p)
    elif suffix in (".txt", ".md"):
        return p.read_text(encoding="utf-8", errors="ignore")
    elif suffix == ".xlsx":
        return _parse_xlsx(p)
    else:
        raise ValueError(f"不支持的文件类型: {suffix}")


def _parse_pdf(p: Path) -> str:
    from pypdf import PdfReader
    reader = PdfReader(str(p))
    return "\n".join(page.extract_text() or "" for page in reader.pages)


def _parse_docx(p: Path) -> str:
    from docx import Document
    doc = Document(str(p))
    return "\n".join(para.text for para in doc.paragraphs)


def _parse_xlsx(p: Path) -> str:
    from openpyxl import load_workbook
    wb = load_workbook(str(p), data_only=True)
    lines = []
    for sheet in wb.worksheets:
        lines.append(f"# Sheet: {sheet.title}")
        for row in sheet.iter_rows(values_only=True):
            cells = [str(c) if c is not None else "" for c in row]
            lines.append("\t".join(cells))
    return "\n".join(lines)